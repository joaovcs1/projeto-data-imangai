"""Modulo 2 - Analise de emissao de boletos.

Entradas:s
  - planilha_limpa.xlsx (aba Boletos; coluna Cliente)
  - base_geral.xlsm (ou base_geral*.xlsm; aba Planilha de boletos)
Saida: 
  - planilha_inadimplencia.xlsx (Analise de emissao + Historico completo).
"""

from pathlib import Path
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

PASTA = Path(__file__).resolve().parent
PLANILHA_LIMPA = PASTA / "planilha_limpa.xlsx"
SAIDA = PASTA / "planilha_inadimplencia.xlsx"
ABA_HISTORICO_ORIGINAL = "Planilha de boletos"

AZUL_ESCURO = "183153"
AZUL_CLARO = "DCEBFA"
VERDE = "DDF2E2"
VERMELHO = "FCE0E0"
AMARELO = "FFF0BC"
CINZA = "E9EDF1"
BRANCO = "FFFFFF"


def normalizar_nome(valor):
    """Ignora acentos, caixa, espacos repetidos e anotacoes entre parenteses."""
    if pd.isna(valor):
        return ""
    nome = re.sub(r"\([^)]*\)", " ", str(valor))
    nome = " ".join(nome.split()).casefold()
    return "".join(c for c in unicodedata.normalize("NFKD", nome)
                   if not unicodedata.combining(c))


def texto_nome(valor):
    return " ".join(re.sub(r"\([^)]*\)", " ", str(valor)).split())


def normalizar_status(valor):
    if pd.isna(valor):
        return ""
    return " ".join(str(valor).upper().split())


def encontrar_base():
    exato = PASTA / "base_geral.xlsm"
    if exato.exists():
        return exato
    alternativas = sorted(PASTA.glob("base_geral*.xlsm"))
    if len(alternativas) == 1:
        return alternativas[0]
    if not alternativas:
        raise FileNotFoundError("Nao foi encontrada base_geral.xlsm na pasta do script.")
    raise ValueError("Ha varias base_geral*.xlsm. Renomeie a desejada para base_geral.xlsm.")


def cor_status(status):
    if status == "PAGO":
        return VERDE
    if status in {"VENCIDO", "PROTESTADO"}:
        return VERMELHO
    if status == "EM ABERTO":
        return AMARELO
    return CINZA


def preparar_entradas():
    if not PLANILHA_LIMPA.exists():
        raise FileNotFoundError(f"Nao encontrada: {PLANILHA_LIMPA.name}")
    origem = pd.read_excel(PLANILHA_LIMPA, sheet_name="Boletos", dtype={"Cliente": "string"})
    if "Cliente" not in origem.columns:
        raise ValueError("A planilha limpa deve conter a coluna 'Cliente'.")
    origem["_nome"] = origem["Cliente"].map(normalizar_nome)
    origem = origem.loc[origem["_nome"] != ""]
    clientes = (origem.drop_duplicates("_nome")[["Cliente", "_nome"]]
                .sort_values("_nome", kind="stable"))

    base = encontrar_base()
    # A linha 3 tem os titulos; linhas 4 e 5 sao auxiliares e serao filtradas.
    df = pd.read_excel(base, sheet_name=ABA_HISTORICO_ORIGINAL,
                       header=2, dtype=object, engine="openpyxl")
    # Por posicao, para preservar as 20 colunas do historico mesmo sem titulo.
    if df.shape[1] < 16:
        raise ValueError("Estrutura inesperada: menos de 16 colunas na base geral.")
    df["_nome"] = df.iloc[:, 2].map(normalizar_nome)
    df["_venc"] = pd.to_datetime(df.iloc[:, 5], dayfirst=True, errors="coerce")
    df["_status"] = df.iloc[:, 15].map(normalizar_status)
    df["_linha_base"] = df.index + 4  # linha original do Excel
    df = df[df["_nome"] != ""].copy()
    # Estavel em caso de vencimentos iguais: mantem ordem de origem.
    df.sort_values(["_nome", "_venc", "_linha_base"],
                   ascending=[True, False, True], kind="stable", inplace=True,
                   na_position="last")
    return clientes, df, list(df.columns[:-4]), base


def adicionar_cabecalho(ws, valores):
    ws.append(valores)
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        cell.font = Font(color=BRANCO, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 31


def exportar(clientes, df, colunas_originais, arquivo_base):
    wb = Workbook()
    analise = wb.active
    analise.title = "Analise de emissao"
    historico = wb.create_sheet("Historico completo")

    colunas_analise = ["Proprietário", "1º venc.", "1º status", "2º venc.",
                       "2º status", "3º venc.", "3º status", "Histórico", "Resultado"]
    adicionar_cabecalho(analise, colunas_analise)
    # Na aba de historico mantemos todas as colunas e adicionamos a marca de selecao.
    titulos = [str(v) if not str(v).startswith("Unnamed:") else ""
               for v in colunas_originais]
    adicionar_cabecalho(historico, ["Considerado na análise"] + titulos)

    grupos = {chave: grupo for chave, grupo in df.groupby("_nome", sort=False)}
    total_emitir = total_nao = sem_historico = 0

    for _, cliente in clientes.iterrows():
        nome = texto_nome(cliente["Cliente"])
        chave = cliente["_nome"]
        grupo = grupos.get(chave)
        if grupo is None:
            grupo = df.iloc[0:0]
        validos = grupo.loc[grupo["_venc"].notna()]
        recentes = validos.head(3)
        existe_historico = not grupo.empty
        resultado = "EMITIR" if recentes.empty or (recentes["_status"] == "PAGO").any() else "NÃO EMITIR"
        if not existe_historico:
            sem_historico += 1
        if resultado == "EMITIR":
            total_emitir += 1
        else:
            total_nao += 1
        linha = [nome]
        for _, boleto in recentes.iterrows():
            linha.extend([boleto["_venc"].to_pydatetime(), boleto["_status"] or "(VAZIO)"])
        for _ in range(3 - len(recentes)):
            linha.extend([None, None])
        situacao = "SEM HISTÓRICO" if not existe_historico else (
            "SEM DATA VÁLIDA" if validos.empty else f"{len(recentes)} analisado(s)")
        linha.extend([situacao, resultado])
        analise.append(linha)
        ra = analise.max_row
        for coluna_data in (2, 4, 6):
            analise.cell(ra, coluna_data).number_format = "DD/MM/YYYY"
        for coluna_status in (3, 5, 7):
            cell = analise.cell(ra, coluna_status)
            if cell.value:
                cell.fill = PatternFill("solid", fgColor=cor_status(str(cell.value)))
        cell_result = analise.cell(ra, 9)
        cell_result.fill = PatternFill("solid", fgColor=VERDE if resultado == "EMITIR" else VERMELHO)
        cell_result.font = Font(bold=True)
        if not existe_historico:
            continue

        # Linha de separacao com o nome antes de cada bloco.
        historico.append([nome.upper()])
        rh = historico.max_row
        for col in range(1, len(colunas_originais) + 2):
            cell = historico.cell(rh, col)
            cell.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
            cell.font = Font(color=BRANCO, bold=True)
        historico.row_dimensions[rh].height = 23

        indices_recentes = set(recentes.index)
        for indice, boleto in grupo.iterrows():
            incluido = indice in indices_recentes
            dados = ["SIM" if incluido else ""]
            for col in colunas_originais:
                valor = boleto[col]
                dados.append(None if pd.isna(valor) else valor)
            historico.append(dados)
            rh = historico.max_row
            if incluido:
                # Azul na coluna de controle; nao encobre as cores de status.
                historico.cell(rh, 1).fill = PatternFill("solid", fgColor=AZUL_CLARO)
                historico.cell(rh, 1).font = Font(bold=True, color=AZUL_ESCURO)
                historico.cell(rh, 1).border = Border(left=Side(style="medium", color="2563A6"))
            # Status original esta na coluna P; com coluna de controle vira Q (17).
            status_cell = historico.cell(rh, 17)
            status_cell.fill = PatternFill("solid", fgColor=cor_status(boleto["_status"]))
            if incluido:
                status_cell.font = Font(bold=True)
            for col_data in (5, 7, 18):
                historico.cell(rh, col_data).number_format = "DD/MM/YYYY"
        historico.append([])

    analise.freeze_panes = "B2"
    analise.auto_filter.ref = analise.dimensions
    historico.freeze_panes = "D2"
    for col, width in {"A": 34, "B": 17, "C": 19, "D": 17, "E": 19,
                       "F": 17, "G": 19, "H": 22, "I": 19}.items():
        analise.column_dimensions[col].width = width
    for idx in range(1, len(colunas_originais) + 2):
        historico.column_dimensions[get_column_letter(idx)].width = (
            26 if idx == 1 else 23 if idx in (4, 17, 19, 20) else 17)
    historico.column_dimensions["D"].width = 34
    historico.column_dimensions["S"].width = 55
    historico.column_dimensions["T"].width = 45

    wb.save(SAIDA)
    print(f"Base: {arquivo_base.name}")
    print(f"Proprietarios analisados: {len(clientes)}")
    print(f"EMITIR: {total_emitir} | NÃO EMITIR: {total_nao} | Sem historico: {sem_historico}")
    print(f"Arquivo gerado: {SAIDA}")


def main():
    clientes, historico, colunas, base = preparar_entradas()
    exportar(clientes, historico, colunas, base)


if __name__ == "__main__":
    main()
